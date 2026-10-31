"""Shared Excel styling used by every report builder."""
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

MONEY_FMT = '"$"#,##0.00'
MONEY_FMT_PRECISE = '"$"#,##0.0000'
HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
TITLE_FONT = Font(bold=True, size=14)
NOTE_FONT = Font(italic=True, color="808080")
_thin_side = Side(style="thin", color="D9D9D9")
BORDER = Border(left=_thin_side, right=_thin_side, top=_thin_side, bottom=_thin_side)


def style_header_row(ws, row: int, ncols: int):
    """Apply the shared header look to a row of column headers."""
    for col in range(1, ncols + 1):
        c = ws.cell(row=row, column=col)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.border = BORDER
        c.alignment = Alignment(horizontal="center")
