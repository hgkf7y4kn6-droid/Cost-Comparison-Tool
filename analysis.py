"""Spend-history analysis: parsing, savings math, the Excel workbook and the Word executive summary."""
import io
import os
from datetime import datetime
from typing import Any, Dict

import openpyxl
import pandas as pd
from docx import Document
from fastapi import HTTPException
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font

from config import OFFICEBASICS_PORTAL_URL, REPORTS_DIR, VALID_TIERS
from engine import engine, resolve_tiers
from excel_style import MONEY_FMT, NOTE_FONT, TITLE_FONT, style_header_row

def parse_spend_file(content: bytes, ext: str) -> pd.DataFrame:
    if ext == "csv":
        df = pd.read_csv(io.BytesIO(content))
    elif ext in {"xlsx", "xls"}:
        df = pd.read_excel(io.BytesIO(content))
    else:
        raise HTTPException(400, "Spend history file must be .csv, .xlsx, or .xls")

    df.columns = [c.strip().lower() for c in df.columns]
    required = {"date", "location", "product", "quantity"}
    missing = required - set(df.columns)
    if missing:
        raise HTTPException(
            422,
            f"Spend file missing columns: {sorted(missing)}. "
            "Required: date, location, product, quantity, and either unit_price or total_price. "
            "Optional: competitor_uses_per_case.",
        )
    if "unit_price" not in df.columns and "total_price" not in df.columns:
        raise HTTPException(422, "Spend file must include either 'unit_price' or 'total_price'.")
    return df


def analyze_spend(df: pd.DataFrame, recommended_tier: str = "best") -> Dict[str, Any]:
    if recommended_tier not in VALID_TIERS:
        recommended_tier = "best"

    match_cache: Dict[str, Dict[str, Any]] = {}
    line_rows = []

    for _, row in df.iterrows():
        product_desc = str(row["product"])
        quantity = float(row["quantity"])

        if "total_price" in df.columns and pd.notna(row.get("total_price")):
            current_spend = float(row["total_price"])
        else:
            unit_price = float(row["unit_price"])
            current_spend = unit_price * quantity

        comp_uses_per_case = None
        if "competitor_uses_per_case" in df.columns and pd.notna(row.get("competitor_uses_per_case")):
            comp_uses_per_case = float(row["competitor_uses_per_case"])

        try:
            month_label = pd.to_datetime(row["date"]).strftime("%Y-%m")
        except Exception:
            month_label = str(row["date"])

        location = str(row["location"])

        if product_desc not in match_cache:
            vec = engine.embed_text(product_desc[:800])
            match_cache[product_desc] = resolve_tiers(vec)
        match = match_cache[product_desc]

        rec_item = match["tiers"].get(recommended_tier)
        if not rec_item:
            for fallback in ["best", "better", "good"]:
                if match["tiers"].get(fallback):
                    rec_item = match["tiers"][fallback]
                    break

        if not rec_item:
            line_rows.append(
                {
                    "month": month_label,
                    "location": location,
                    "product": product_desc,
                    "quantity": quantity,
                    "current_spend": current_spend,
                    "matched_category": None,
                    "recommended_item": None,
                    "recommended_spend": current_spend,
                    "savings": 0.0,
                    "matched": False,
                    "needs_manual_lookup": True,
                }
            )
            continue

        rec_uses_per_case = rec_item.get("uses_per_case")
        if comp_uses_per_case and rec_uses_per_case:
            total_uses = quantity * comp_uses_per_case
            recommended_cases = total_uses / rec_uses_per_case
            recommended_spend = recommended_cases * rec_item["price"]
        else:
            # No usage normalization data available: assume the same purchase
            # volume (case-for-case substitution). Directional, not exact.
            recommended_spend = quantity * rec_item["price"]

        line_rows.append(
            {
                "month": month_label,
                "location": location,
                "product": product_desc,
                "quantity": quantity,
                "current_spend": current_spend,
                "matched_category": match["matched_category"],
                "recommended_item": f"{rec_item['name']} ({rec_item['sku']})",
                "recommended_spend": recommended_spend,
                "savings": current_spend - recommended_spend,
                "matched": True,
                "needs_manual_lookup": False,
            }
        )

    detail_df = pd.DataFrame(line_rows)
    rows_unmatched = int((~detail_df["matched"]).sum())

    total_current = float(detail_df["current_spend"].sum())
    total_recommended = float(detail_df["recommended_spend"].sum())
    total_savings = total_current - total_recommended
    savings_pct = (total_savings / total_current * 100) if total_current else 0.0

    by_month = (
        detail_df.groupby("month")
        .agg(current_spend=("current_spend", "sum"), recommended_spend=("recommended_spend", "sum"))
        .reset_index()
    )
    by_month["savings"] = by_month["current_spend"] - by_month["recommended_spend"]
    by_month = by_month.sort_values("month")

    by_location = (
        detail_df.groupby("location")
        .agg(current_spend=("current_spend", "sum"), recommended_spend=("recommended_spend", "sum"))
        .reset_index()
    )
    by_location["savings"] = by_location["current_spend"] - by_location["recommended_spend"]
    by_location = by_location.sort_values("savings", ascending=False)

    matched_df = detail_df[detail_df["matched"]]
    by_category = (
        matched_df.groupby("matched_category")
        .agg(current_spend=("current_spend", "sum"), recommended_spend=("recommended_spend", "sum"))
        .reset_index()
    )
    by_category["savings"] = by_category["current_spend"] - by_category["recommended_spend"]
    by_category = by_category.sort_values("savings", ascending=False)

    return {
        "recommended_tier": recommended_tier,
        "total_current_spend": total_current,
        "total_recommended_spend": total_recommended,
        "total_savings": total_savings,
        "savings_pct": savings_pct,
        "rows_analyzed": len(detail_df),
        "rows_unmatched": rows_unmatched,
        "manual_lookup_note": (
            f"{rows_unmatched} record(s) had no catalog match. supplies.officebasics.com "
            "requires an account login and can't be searched automatically - check those "
            "products manually there until more catalog items are uploaded."
            if rows_unmatched
            else None
        ),
        "officebasics_portal_url": OFFICEBASICS_PORTAL_URL if rows_unmatched else None,
        "by_month": by_month.to_dict(orient="records"),
        "by_location": by_location.to_dict(orient="records"),
        "by_category": by_category.to_dict(orient="records"),
        "detail": detail_df.drop(columns=["matched"]).to_dict(orient="records"),
    }


def build_spend_analysis_workbook(analysis: Dict[str, Any]) -> str:
    money_fmt = MONEY_FMT  # local alias kept for readability in this function
    wb = openpyxl.Workbook()

    # --- Summary ---
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "Office Basics Spend Analysis"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = (
        f"Generated {datetime.now().strftime('%B %d, %Y')} \u2014 "
        f"recommended tier: {analysis['recommended_tier'].upper()}"
    )
    ws["A2"].font = NOTE_FONT

    summary_rows = [
        ("Total Current Spend", analysis["total_current_spend"], money_fmt),
        ("Total Recommended Spend", analysis["total_recommended_spend"], money_fmt),
        ("Total Savings", analysis["total_savings"], money_fmt),
        ("Savings %", analysis["savings_pct"] / 100, "0.0%"),
        ("Records Analyzed", analysis["rows_analyzed"], None),
        ("Records Unmatched", analysis["rows_unmatched"], None),
    ]
    for i, (label, value, fmt) in enumerate(summary_rows, start=4):
        ws.cell(row=i, column=1, value=label).font = Font(bold=True)
        c = ws.cell(row=i, column=2, value=value)
        if fmt:
            c.number_format = fmt
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 20

    # --- By Month ---
    ws2 = wb.create_sheet("By Month")
    headers = ["Month", "Current Spend", "Recommended Spend", "Savings"]
    for col, h in enumerate(headers, start=1):
        ws2.cell(row=1, column=col, value=h)
    style_header_row(ws2, 1, len(headers))
    for i, row in enumerate(analysis["by_month"], start=2):
        ws2.cell(row=i, column=1, value=row["month"])
        ws2.cell(row=i, column=2, value=row["current_spend"]).number_format = money_fmt
        ws2.cell(row=i, column=3, value=row["recommended_spend"]).number_format = money_fmt
        ws2.cell(row=i, column=4, value=row["savings"]).number_format = money_fmt
    for col, w in zip("ABCD", [14, 18, 20, 14]):
        ws2.column_dimensions[col].width = w
    if analysis["by_month"]:
        n = len(analysis["by_month"])
        chart = BarChart()
        chart.title = "Current vs. Recommended Spend by Month"
        chart.y_axis.title = "Spend ($)"
        data = Reference(ws2, min_col=2, max_col=3, min_row=1, max_row=n + 1)
        cats = Reference(ws2, min_col=1, min_row=2, max_row=n + 1)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        ws2.add_chart(chart, "F2")

    # --- By Location ---
    ws3 = wb.create_sheet("By Location")
    headers_loc = ["Location", "Current Spend", "Recommended Spend", "Savings"]
    for col, h in enumerate(headers_loc, start=1):
        ws3.cell(row=1, column=col, value=h)
    style_header_row(ws3, 1, len(headers_loc))
    for i, row in enumerate(analysis["by_location"], start=2):
        ws3.cell(row=i, column=1, value=row["location"])
        ws3.cell(row=i, column=2, value=row["current_spend"]).number_format = money_fmt
        ws3.cell(row=i, column=3, value=row["recommended_spend"]).number_format = money_fmt
        ws3.cell(row=i, column=4, value=row["savings"]).number_format = money_fmt
    for col, w in zip("ABCD", [22, 18, 20, 14]):
        ws3.column_dimensions[col].width = w
    if analysis["by_location"]:
        n = len(analysis["by_location"])
        chart2 = BarChart()
        chart2.title = "Savings Opportunity by Location"
        chart2.y_axis.title = "Spend ($)"
        data = Reference(ws3, min_col=2, max_col=3, min_row=1, max_row=n + 1)
        cats = Reference(ws3, min_col=1, min_row=2, max_row=n + 1)
        chart2.add_data(data, titles_from_data=True)
        chart2.set_categories(cats)
        ws3.add_chart(chart2, "F2")

    # --- By Category ---
    ws4 = wb.create_sheet("By Category")
    headers_cat = ["Category", "Current Spend", "Recommended Spend", "Savings"]
    for col, h in enumerate(headers_cat, start=1):
        ws4.cell(row=1, column=col, value=h)
    style_header_row(ws4, 1, len(headers_cat))
    for i, row in enumerate(analysis["by_category"], start=2):
        ws4.cell(row=i, column=1, value=row["matched_category"])
        ws4.cell(row=i, column=2, value=row["current_spend"]).number_format = money_fmt
        ws4.cell(row=i, column=3, value=row["recommended_spend"]).number_format = money_fmt
        ws4.cell(row=i, column=4, value=row["savings"]).number_format = money_fmt
    for col, w in zip("ABCD", [26, 18, 20, 14]):
        ws4.column_dimensions[col].width = w

    # --- Line Detail ---
    ws5 = wb.create_sheet("Line Detail")
    headers_d = [
        "Month", "Location", "Product", "Qty", "Current Spend",
        "Matched Category", "Recommended Item", "Recommended Spend", "Savings",
    ]
    for col, h in enumerate(headers_d, start=1):
        ws5.cell(row=1, column=col, value=h)
    style_header_row(ws5, 1, len(headers_d))
    for i, row in enumerate(analysis["detail"], start=2):
        ws5.cell(row=i, column=1, value=row["month"])
        ws5.cell(row=i, column=2, value=row["location"])
        ws5.cell(row=i, column=3, value=row["product"])
        ws5.cell(row=i, column=4, value=row["quantity"])
        ws5.cell(row=i, column=5, value=row["current_spend"]).number_format = money_fmt
        ws5.cell(row=i, column=6, value=row["matched_category"])
        ws5.cell(row=i, column=7, value=row["recommended_item"])
        ws5.cell(row=i, column=8, value=row["recommended_spend"]).number_format = money_fmt
        ws5.cell(row=i, column=9, value=row["savings"]).number_format = money_fmt
    for col_letter, w in zip("ABCDEFGHI", [12, 18, 30, 8, 14, 20, 30, 16, 12]):
        ws5.column_dimensions[col_letter].width = w

    filename = f"spend_analysis_{int(datetime.now().timestamp())}.xlsx"
    wb.save(os.path.join(REPORTS_DIR, filename))
    return filename


def build_executive_summary_docx(analysis: Dict[str, Any]) -> str:
    doc = Document()

    doc.add_heading("Office Basics Spend Analysis \u2014 Executive Summary", level=0)
    sub = doc.add_paragraph(f"Generated {datetime.now().strftime('%B %d, %Y')}")
    sub.runs[0].italic = True

    n_months = len(analysis["by_month"])
    n_locations = len(analysis["by_location"])
    doc.add_paragraph(
        f"This analysis reviewed {analysis['rows_analyzed']} purchase records across "
        f"{n_locations} location(s) and {n_months} month(s), totaling "
        f"${analysis['total_current_spend']:,.2f} in tracked spend. Based on "
        f"{analysis['recommended_tier'].upper()}-tier Office Basics equivalents, an estimated "
        f"${analysis['total_savings']:,.2f} ({analysis['savings_pct']:.1f}%) could have been saved "
        "over this period."
    )

    doc.add_heading("Key Figures", level=1)
    table = doc.add_table(rows=1, cols=2)
    table.style = "Light List Accent 1"
    hdr = table.rows[0].cells
    hdr[0].text, hdr[1].text = "Metric", "Value"
    for label, value in [
        ("Total Current Spend", f"${analysis['total_current_spend']:,.2f}"),
        ("Total Recommended Spend", f"${analysis['total_recommended_spend']:,.2f}"),
        ("Total Savings", f"${analysis['total_savings']:,.2f}"),
        ("Savings %", f"{analysis['savings_pct']:.1f}%"),
        ("Records Analyzed", str(analysis["rows_analyzed"])),
        ("Records Unmatched", str(analysis["rows_unmatched"])),
    ]:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value

    if analysis["by_category"]:
        doc.add_heading("Top Savings Opportunities by Category", level=1)
        top_categories = sorted(analysis["by_category"], key=lambda r: r["savings"], reverse=True)[:5]
        cat_table = doc.add_table(rows=1, cols=3)
        cat_table.style = "Light List Accent 1"
        hdr = cat_table.rows[0].cells
        hdr[0].text, hdr[1].text, hdr[2].text = "Category", "Current Spend", "Savings"
        for row in top_categories:
            cells = cat_table.add_row().cells
            cells[0].text = str(row["matched_category"])
            cells[1].text = f"${row['current_spend']:,.2f}"
            cells[2].text = f"${row['savings']:,.2f}"

    if analysis["by_location"]:
        doc.add_heading("Spend & Savings by Location", level=1)
        loc_table = doc.add_table(rows=1, cols=4)
        loc_table.style = "Light List Accent 1"
        hdr = loc_table.rows[0].cells
        hdr[0].text, hdr[1].text, hdr[2].text, hdr[3].text = (
            "Location", "Current Spend", "Recommended Spend", "Savings",
        )
        for row in analysis["by_location"]:
            cells = loc_table.add_row().cells
            cells[0].text = str(row["location"])
            cells[1].text = f"${row['current_spend']:,.2f}"
            cells[2].text = f"${row['recommended_spend']:,.2f}"
            cells[3].text = f"${row['savings']:,.2f}"

    doc.add_heading("Methodology & Assumptions", level=1)
    for bullet in [
        f"Recommended spend uses the {analysis['recommended_tier'].upper()} tier Office Basics "
        "equivalent for each matched product.",
        "Where per-case usage data (uses_per_case) is available for both the current product and "
        "its recommended replacement, spend is normalized to cost-per-use before comparison. "
        "Otherwise, a case-for-case substitution is assumed \u2014 treat those figures as directional.",
        f"{analysis['rows_unmatched']} of {analysis['rows_analyzed']} records could not be matched "
        "to a catalog category and are included in current spend totals with no assumed savings.",
        "Product matching uses a combination of previously confirmed recommendations and "
        "text/visual similarity against the loaded Office Basics catalog.",
    ]:
        doc.add_paragraph(bullet, style="List Bullet")

    if analysis.get("manual_lookup_note"):
        doc.add_paragraph(
            f"{analysis['manual_lookup_note']} Portal: {analysis['officebasics_portal_url']}",
            style="List Bullet",
        )

    doc.add_heading("Recommendation", level=1)
    doc.add_paragraph(
        "Prioritize rolling out Office Basics equivalents in the highest-savings categories and "
        "locations identified above, where the switching cost is lowest and the dollar impact is "
        "greatest."
    )

    filename = f"executive_summary_{int(datetime.now().timestamp())}.docx"
    doc.save(os.path.join(REPORTS_DIR, filename))
    return filename
