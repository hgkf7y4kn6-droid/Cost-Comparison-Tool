import io
import json
import os
import re
import tempfile
import uuid
from datetime import datetime
from typing import Any, Dict, List, Optional

import faiss
import numpy as np
import openpyxl
import pandas as pd
import pdfplumber
import torch
from docx import Document
from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from PIL import Image
from pydantic import BaseModel
from transformers import CLIPModel, CLIPProcessor

try:
    import pytesseract

    OCR_AVAILABLE = True
except ImportError:
    OCR_AVAILABLE = False

VALID_TIERS = {"good", "better", "best"}
OFFICEBASICS_PORTAL_URL = "https://supplies.officebasics.com/"
# DATA_DIR lets a host (e.g. a Render persistent disk) keep reports and learned
# feedback across restarts; defaults to the system temp dir for local runs.
REPORTS_DIR = os.path.join(os.environ.get("DATA_DIR", tempfile.gettempdir()), "ob_reports")
os.makedirs(REPORTS_DIR, exist_ok=True)
FEEDBACK_PATH = os.path.join(REPORTS_DIR, "feedback_store.json")
FEEDBACK_SIMILARITY_THRESHOLD = 0.88  # cosine sim (CLIP embeddings are L2-normalized)


# ==========================================
# 1. CLIP + FAISS ENGINE
# ==========================================
class CLIPSearchEngine:
    def __init__(self, model_name: str = "openai/clip-vit-base-patch32"):
        self.device = "cuda" if torch.cuda.is_available() else "cpu"
        self.model = CLIPModel.from_pretrained(model_name).to(self.device)
        self.processor = CLIPProcessor.from_pretrained(model_name)
        self.dimension = 512
        self.index = faiss.IndexFlatIP(self.dimension)
        self.catalog: List[Dict[str, Any]] = []
        self.by_sku: Dict[str, Dict[str, Any]] = {}

    def embed_image(self, image: Image.Image) -> np.ndarray:
        inputs = self.processor(images=image, return_tensors="pt").to(self.device)
        with torch.no_grad():
            features = self.model.get_image_features(**inputs)
        vec = features.cpu().numpy().astype("float32")
        faiss.normalize_L2(vec)
        return vec

    def embed_text(self, text: str) -> np.ndarray:
        inputs = self.processor(
            text=[text], return_tensors="pt", padding=True, truncation=True
        ).to(self.device)
        with torch.no_grad():
            features = self.model.get_text_features(**inputs)
        vec = features.cpu().numpy().astype("float32")
        faiss.normalize_L2(vec)
        return vec

    def load_catalog(self, items: List[Dict[str, Any]]):
        """Each item must have: sku, name, description, category, tier, price.
        Optional: format, uses_per_case."""
        self.catalog = items
        self.by_sku = {it["sku"]: it for it in items}
        self.index = faiss.IndexFlatIP(self.dimension)
        if not items:
            return
        vectors = [
            self.embed_text(f"{it['name']} - {it['description']}")[0] for it in items
        ]
        self.index.add(np.array(vectors).astype("float32"))

    def search(self, vector: np.ndarray, top_k: int = 30) -> List[Dict[str, Any]]:
        if self.index.ntotal == 0:
            return []
        top_k = min(top_k, self.index.ntotal)
        scores, indices = self.index.search(vector, top_k)
        results = []
        for score, idx in zip(scores[0], indices[0]):
            if idx != -1:
                item = self.catalog[idx].copy()
                item["similarity_score"] = float(score)
                results.append(item)
        return results


# ==========================================
# 2. LEARNED FEEDBACK STORE (RAG layer)
# ==========================================
class FeedbackStore:
    """
    Remembers which SKUs a rep confirmed as the good/better/best match for a
    given competitor product. Future uploads of a visually/textually similar
    competitor product retrieve this confirmed mapping instead of recomputing
    from scratch.
    """

    def __init__(self, engine: CLIPSearchEngine, path: str = FEEDBACK_PATH):
        self.engine = engine
        self.path = path
        self.entries: List[Dict[str, Any]] = []
        self.index = faiss.IndexFlatIP(engine.dimension)
        self._load()

    def _load(self):
        if os.path.exists(self.path):
            try:
                with open(self.path) as f:
                    self.entries = json.load(f)
            except (json.JSONDecodeError, OSError):
                self.entries = []
        self.index = faiss.IndexFlatIP(self.engine.dimension)
        if self.entries:
            vectors = np.array([e["vector"] for e in self.entries]).astype("float32")
            self.index.add(vectors)

    def _save(self):
        with open(self.path, "w") as f:
            json.dump(self.entries, f)

    def add(self, vector: np.ndarray, competitor_name: str, tier_skus: Dict[str, str]):
        entry = {
            "vector": vector[0].tolist(),
            "competitor_name": competitor_name,
            "tier_skus": tier_skus,
            "confirmed_at": datetime.now().isoformat(),
        }
        self.entries.append(entry)
        self.index.add(vector.astype("float32"))
        self._save()

    def lookup(self, vector: np.ndarray, threshold: float = FEEDBACK_SIMILARITY_THRESHOLD) -> Optional[Dict[str, Any]]:
        if self.index.ntotal == 0:
            return None
        scores, indices = self.index.search(vector.astype("float32"), 1)
        score, idx = float(scores[0][0]), int(indices[0][0])
        if idx == -1 or score < threshold:
            return None
        return {"match_score": score, **self.entries[idx]}


# ==========================================
# 3. TEXT / PRICE EXTRACTION
# ==========================================
PRICE_PATTERN = re.compile(r"\$\s?(\d{1,3}(?:,\d{3})*(?:\.\d{1,2})?)")
PRICE_KEYWORDS = ["price", "cost", "total", "unit price", "each", "$/unit", "list price"]


def extract_price(text: str) -> Optional[float]:
    matches = PRICE_PATTERN.findall(text)
    if not matches:
        return None
    return float(matches[0].replace(",", ""))


def extract_price_near_keywords(text: str) -> Optional[float]:
    for line in text.splitlines():
        if any(k in line.lower() for k in PRICE_KEYWORDS):
            price = extract_price(line)
            if price is not None:
                return price
    return extract_price(text)


def extract_text_and_image(content: bytes, ext: str):
    text = ""
    image = None

    if ext in {"png", "jpg", "jpeg", "webp"}:
        image = Image.open(io.BytesIO(content)).convert("RGB")
        if OCR_AVAILABLE:
            try:
                text = pytesseract.image_to_string(image)
            except Exception:
                text = ""

    elif ext == "pdf":
        with pdfplumber.open(io.BytesIO(content)) as pdf:
            text = " ".join(p.extract_text() or "" for p in pdf.pages)
            if not text.strip() and OCR_AVAILABLE and pdf.pages:
                rendered = pdf.pages[0].to_image(resolution=200).original
                image = rendered.convert("RGB")
                text = pytesseract.image_to_string(image)

    elif ext in {"xlsx", "xls"}:
        df = pd.read_excel(io.BytesIO(content))
        text = " ".join(df.fillna("").astype(str).values.flatten())

    else:
        raise HTTPException(status_code=400, detail=f"Unsupported file format: .{ext}")

    return text, image


def cost_per_use(price: Optional[float], uses_per_case: Optional[float]) -> Optional[float]:
    if price is None or not uses_per_case:
        return None
    return price / uses_per_case


# ==========================================
# Shared Excel styling (used by every report builder)
# ==========================================
MONEY_FMT = '"$"#,##0.00'
MONEY_FMT_PRECISE = '"$"#,##0.0000'
HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
TITLE_FONT = Font(bold=True, size=14)
NOTE_FONT = Font(italic=True, color="808080")
_thin_side = Side(style="thin", color="D9D9D9")
BORDER = Border(left=_thin_side, right=_thin_side, top=_thin_side, bottom=_thin_side)


def style_header_row(ws, row: int, ncols: int):
    """Apply the shared header look to a row of column headers."""
    for col in range(1, ncols + 1):
        c = ws.cell(row=row, column=col)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.border = BORDER
        c.alignment = Alignment(horizontal="center")


# ==========================================
# 4. API + shared engine/store
# ==========================================
app = FastAPI(title="Office Basics Good/Better/Best Comparison Engine")

# The frontend is served from a different origin (Cloudflare), so the browser
# needs CORS. ALLOWED_ORIGINS is a comma-separated list, e.g.
# "https://cost-comparison-tool.example.workers.dev"; "*" allows any origin.
ALLOWED_ORIGINS = [o.strip() for o in os.environ.get("ALLOWED_ORIGINS", "*").split(",") if o.strip()]
app.add_middleware(
    CORSMiddleware,
    allow_origins=ALLOWED_ORIGINS,
    allow_methods=["*"],
    allow_headers=["*"],
)
app.mount("/reports", StaticFiles(directory=REPORTS_DIR), name="reports")
engine = CLIPSearchEngine()
feedback_store = FeedbackStore(engine)

# in-memory cache of pending (unconfirmed) comparisons, keyed by comparison_id,
# so /feedback/confirm/ doesn't require re-uploading the file. Bounded so a
# long-running process can't accumulate this without limit - old, never-
# confirmed comparisons are simply no longer confirmable, which only means
# the rep has to re-run /compare/ before confirming.
PENDING_COMPARISONS_MAX = 2000
PENDING_COMPARISONS: Dict[str, Dict[str, Any]] = {}


def _remember_pending_comparison(comparison_id: str, payload: Dict[str, Any]) -> None:
    if len(PENDING_COMPARISONS) >= PENDING_COMPARISONS_MAX:
        oldest_id = next(iter(PENDING_COMPARISONS))
        del PENDING_COMPARISONS[oldest_id]
    PENDING_COMPARISONS[comparison_id] = payload


def resolve_tiers(vec: np.ndarray) -> Dict[str, Any]:
    """
    Shared matching core used by both single-item comparisons and the spend
    analysis: checks the learned feedback store first, then falls back to
    algorithmic category + per-tier similarity ranking for anything missing.
    """
    tiers: Dict[str, Dict[str, Any]] = {}
    source = "algorithmic"
    matched_category = None
    match_score = None

    learned = feedback_store.lookup(vec)
    if learned:
        source = "learned"
        match_score = learned["match_score"]
        for tier_name, sku in learned["tier_skus"].items():
            item = engine.by_sku.get(sku)
            if item:
                tiers[tier_name] = {k: v for k, v in item.items() if k != "similarity_score"}
                matched_category = matched_category or item["category"]

    if len(tiers) < 3:
        matches = engine.search(vec, top_k=30)
        if matches:
            top_category = matched_category or matches[0]["category"]
            same_category = [m for m in matches if m["category"] == top_category]
            for tier_name in ["good", "better", "best"]:
                if tier_name in tiers:
                    continue
                candidates = [m for m in same_category if m["tier"] == tier_name]
                if candidates:
                    best = max(candidates, key=lambda x: x["similarity_score"])
                    tiers[tier_name] = {k: v for k, v in best.items() if k != "similarity_score"}
            matched_category = matched_category or top_category

    return {
        "tiers": tiers,
        "matched_category": matched_category,
        "source": source,
        "match_score": match_score,
    }


def manual_lookup_hint(tiers: Dict[str, Any]) -> Optional[Dict[str, Any]]:
    """
    supplies.officebasics.com is a login-gated B2B portal (per-account contract
    pricing, JS-rendered) and cannot be searched or scraped automatically. When
    a tier is missing from the catalog match, surface a manual-lookup pointer
    instead of silently failing or fabricating a result.
    """
    missing = [t for t in ["good", "better", "best"] if t not in tiers]
    if not missing:
        return None
    return {
        "missing_tiers": missing,
        "note": (
            f"No catalog match yet for: {', '.join(missing)}. supplies.officebasics.com "
            "requires an account login and can't be searched automatically - check it "
            "manually there until more catalog items are uploaded."
        ),
        "portal_url": OFFICEBASICS_PORTAL_URL,
    }


# ==========================================
# 5. CATALOG UPLOAD
# ==========================================
@app.post("/catalog/upload/")
async def upload_catalog(file: UploadFile = File(...)):
    """
    Ingest the Office Basics SKU/price file.
    Required columns (case-insensitive): sku, name, description, category, tier, price
    Optional columns: format, uses_per_case
    tier must be one of: good, better, best (other rows are kept out of matching but not rejected)
    """
    content = await file.read()
    ext = file.filename.lower().split(".")[-1]
    if ext not in {"xlsx", "xls", "csv"}:
        raise HTTPException(400, "Catalog file must be .xlsx, .xls, or .csv")

    df = pd.read_csv(io.BytesIO(content)) if ext == "csv" else pd.read_excel(io.BytesIO(content))
    df.columns = [c.strip().lower() for c in df.columns]

    required = {"sku", "name", "description", "category", "tier", "price"}
    missing = required - set(df.columns)
    if missing:
        raise HTTPException(422, f"Catalog file missing columns: {sorted(missing)}")

    items = []
    skipped = 0
    for _, row in df.iterrows():
        tier = str(row["tier"]).strip().lower()
        if tier not in VALID_TIERS:
            skipped += 1
            continue
        try:
            price = float(row["price"])
        except (ValueError, TypeError):
            skipped += 1
            continue

        uses_per_case = None
        if "uses_per_case" in df.columns:
            try:
                raw = row["uses_per_case"]
                if pd.notna(raw) and str(raw).strip() not in {"", "TBD"}:
                    uses_per_case = float(raw)
            except (ValueError, TypeError):
                uses_per_case = None

        fmt = None
        if "format" in df.columns and pd.notna(row["format"]):
            fmt = str(row["format"]).strip()

        items.append(
            {
                "sku": str(row["sku"]),
                "name": str(row["name"]),
                "description": str(row["description"]),
                "category": str(row["category"]).strip().lower(),
                "tier": tier,
                "price": price,
                "uses_per_case": uses_per_case,
                "format": fmt,
            }
        )

    if not items:
        raise HTTPException(422, "No valid rows found (check tier values are good/better/best and price is numeric).")

    engine.load_catalog(items)
    return {
        "loaded_items": len(items),
        "skipped_rows": skipped,
        "categories": sorted({i["category"] for i in items}),
    }


@app.get("/")
async def root():
    # This is an API only; the UI is the separately hosted frontend.
    return {
        "service": "Office Basics Comparison API",
        "status": "ok",
        "docs": "/docs",
        "health": "/health",
    }


@app.get("/health")
async def health():
    return {"status": "ok"}


@app.get("/catalog/status/")
async def catalog_status():
    return {
        "loaded_items": len(engine.catalog),
        "categories": sorted({i["category"] for i in engine.catalog}),
        "learned_comparisons": len(feedback_store.entries),
        "ocr_available": OCR_AVAILABLE,
    }


# ==========================================
# 6. SINGLE-ITEM COMPARISON + EXCEL REPORT
# ==========================================
def build_comparison_workbook(comparison: Dict[str, Any]) -> str:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Cost Comparison"

    ws["A1"] = "Office Basics Cost Comparison"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:F1")
    source_note = "Learned from a confirmed match" if comparison.get("source") == "learned" else "Algorithmic match"
    ws["A2"] = f"Generated {datetime.now().strftime('%B %d, %Y')} \u2014 {source_note}"
    ws["A2"].font = NOTE_FONT

    headers = ["Item", "SKU", "Description", "Price", "Cost / Use", "Savings vs Current"]
    header_row = 4
    for col, h in enumerate(headers, start=1):
        ws.cell(row=header_row, column=col, value=h)
    style_header_row(ws, header_row, len(headers))

    competitor_price = comparison.get("competitor_price")
    competitor_cpu = comparison.get("competitor_cost_per_use")

    rows = [
        {
            "label": f"CURRENT: {comparison['competitor_product']}",
            "sku": "\u2014",
            "description": "Currently purchased product (competitor)",
            "price": competitor_price,
            "cpu": competitor_cpu,
            "is_option": False,
        }
    ]
    for tier_name in ["good", "better", "best"]:
        item = comparison["tiers"].get(tier_name)
        if item:
            item_cpu = cost_per_use(item.get("price"), item.get("uses_per_case"))
            rows.append(
                {
                    "label": f"{tier_name.upper()}: {item['name']}",
                    "sku": item["sku"],
                    "description": item["description"],
                    "price": item["price"],
                    "cpu": item_cpu,
                    "is_option": True,
                }
            )

    baseline = competitor_cpu if competitor_cpu is not None else competitor_price

    r = header_row + 1
    for row in rows:
        ws.cell(row=r, column=1, value=row["label"]).border = BORDER
        ws.cell(row=r, column=2, value=row["sku"]).border = BORDER
        ws.cell(row=r, column=3, value=row["description"]).border = BORDER

        price_cell = ws.cell(row=r, column=4, value=row["price"])
        price_cell.number_format = MONEY_FMT
        price_cell.border = BORDER

        cpu_cell = ws.cell(row=r, column=5, value=row["cpu"])
        if row["cpu"] is not None:
            cpu_cell.number_format = MONEY_FMT_PRECISE
        else:
            cpu_cell.value = "n/a"
        cpu_cell.border = BORDER

        savings_cell = ws.cell(row=r, column=6)
        savings_cell.border = BORDER
        if row["is_option"] and baseline:
            compare_value = row["cpu"] if competitor_cpu is not None else row["price"]
            if compare_value is not None:
                savings = baseline - compare_value
                pct = (savings / baseline) * 100
                unit_label = "/use" if competitor_cpu is not None else ""
                savings_cell.value = f"${savings:,.4f}{unit_label} ({pct:.1f}%)"
                savings_cell.font = Font(color="1F7A1F" if savings > 0 else "B22222")
        r += 1

    for i, w in enumerate([34, 14, 45, 12, 12, 22], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    filename = f"cost_comparison_{int(datetime.now().timestamp())}.xlsx"
    wb.save(os.path.join(REPORTS_DIR, filename))
    return filename


async def _run_comparison(
    file: UploadFile, competitor_uses_per_case: Optional[float] = None
) -> Dict[str, Any]:
    if not engine.catalog:
        raise HTTPException(400, "No Office Basics catalog loaded. POST one to /catalog/upload/ first.")

    filename = file.filename.lower()
    ext = filename.split(".")[-1]
    content = await file.read()
    text, image = extract_text_and_image(content, ext)

    competitor_price = extract_price_near_keywords(text)

    if image is not None and ext in {"png", "jpg", "jpeg", "webp"}:
        vec = engine.embed_image(image)
    else:
        if not text.strip():
            raise HTTPException(422, "Could not extract any text from the document.")
        vec = engine.embed_text(text[:800])

    competitor_name = filename
    for line in text.splitlines():
        if line.strip():
            competitor_name = line.strip()[:120]
            break

    resolved = resolve_tiers(vec)
    competitor_cpu = cost_per_use(competitor_price, competitor_uses_per_case)

    comparison_id = str(uuid.uuid4())
    result = {
        "comparison_id": comparison_id,
        "competitor_product": competitor_name,
        "competitor_price": competitor_price,
        "competitor_uses_per_case": competitor_uses_per_case,
        "competitor_cost_per_use": competitor_cpu,
        "matched_category": resolved["matched_category"],
        "source": resolved["source"],
        "match_score": resolved["match_score"],
        "tiers": resolved["tiers"],
        "manual_lookup": manual_lookup_hint(resolved["tiers"]),
    }

    _remember_pending_comparison(comparison_id, {
        "vector": vec,
        "competitor_name": competitor_name,
        "suggested_tier_skus": {t: v["sku"] for t, v in resolved["tiers"].items()},
    })

    return result


@app.post("/compare/")
async def compare_product(
    file: UploadFile = File(...),
    competitor_uses_per_case: Optional[float] = Form(None),
):
    """Returns the good/better/best match JSON without generating a file."""
    return await _run_comparison(file, competitor_uses_per_case)


@app.post("/compare/download/")
async def compare_and_download(
    file: UploadFile = File(...),
    competitor_uses_per_case: Optional[float] = Form(None),
):
    """Runs the comparison and generates a downloadable Excel cost-comparison report."""
    comparison = await _run_comparison(file, competitor_uses_per_case)
    filename = build_comparison_workbook(comparison)
    return {"comparison": comparison, "download_url": f"/reports/{filename}"}


class ConfirmFeedbackRequest(BaseModel):
    comparison_id: str
    good_sku: Optional[str] = None
    better_sku: Optional[str] = None
    best_sku: Optional[str] = None


@app.post("/feedback/confirm/")
async def confirm_feedback(payload: ConfirmFeedbackRequest):
    """
    Records which SKUs actually got used for a given competitor product.
    Only the tiers explicitly passed are overridden; anything left blank
    keeps the tier the comparison originally suggested. Future uploads of a
    similar-looking competitor product will retrieve this confirmed mapping.
    """
    pending = PENDING_COMPARISONS.get(payload.comparison_id)
    if not pending:
        raise HTTPException(404, "Unknown or expired comparison_id. Re-run /compare/ first.")

    overrides = {
        "good": payload.good_sku,
        "better": payload.better_sku,
        "best": payload.best_sku,
    }
    final_tier_skus = dict(pending["suggested_tier_skus"])
    for tier, sku in overrides.items():
        if sku:
            if sku not in engine.by_sku:
                raise HTTPException(422, f"SKU '{sku}' not found in loaded catalog.")
            final_tier_skus[tier] = sku

    if not final_tier_skus:
        raise HTTPException(422, "No tier/SKU information available to confirm.")

    feedback_store.add(pending["vector"], pending["competitor_name"], final_tier_skus)
    return {"status": "confirmed", "competitor_name": pending["competitor_name"], "tier_skus": final_tier_skus}


@app.get("/feedback/list/")
async def list_feedback():
    return {
        "count": len(feedback_store.entries),
        "entries": [
            {
                "competitor_name": e["competitor_name"],
                "tier_skus": e["tier_skus"],
                "confirmed_at": e["confirmed_at"],
            }
            for e in feedback_store.entries
        ],
    }


# ==========================================
# 7. SPEND HISTORY / USAGE ANALYSIS
# ==========================================
def parse_spend_file(content: bytes, ext: str) -> pd.DataFrame:
    if ext == "csv":
        df = pd.read_csv(io.BytesIO(content))
    elif ext in {"xlsx", "xls"}:
        df = pd.read_excel(io.BytesIO(content))
    else:
        raise HTTPException(400, "Spend history file must be .csv, .xlsx, or .xls")

    df.columns = [c.strip().lower() for c in df.columns]
    required = {"date", "location", "product", "quantity"}
    missing = required - set(df.columns)
    if missing:
        raise HTTPException(
            422,
            f"Spend file missing columns: {sorted(missing)}. "
            "Required: date, location, product, quantity, and either unit_price or total_price. "
            "Optional: competitor_uses_per_case.",
        )
    if "unit_price" not in df.columns and "total_price" not in df.columns:
        raise HTTPException(422, "Spend file must include either 'unit_price' or 'total_price'.")
    return df


def analyze_spend(df: pd.DataFrame, recommended_tier: str = "best") -> Dict[str, Any]:
    if recommended_tier not in VALID_TIERS:
        recommended_tier = "best"

    match_cache: Dict[str, Dict[str, Any]] = {}
    line_rows = []

    for _, row in df.iterrows():
        product_desc = str(row["product"])
        quantity = float(row["quantity"])

        if "total_price" in df.columns and pd.notna(row.get("total_price")):
            current_spend = float(row["total_price"])
        else:
            unit_price = float(row["unit_price"])
            current_spend = unit_price * quantity

        comp_uses_per_case = None
        if "competitor_uses_per_case" in df.columns and pd.notna(row.get("competitor_uses_per_case")):
            comp_uses_per_case = float(row["competitor_uses_per_case"])

        try:
            month_label = pd.to_datetime(row["date"]).strftime("%Y-%m")
        except Exception:
            month_label = str(row["date"])

        location = str(row["location"])

        if product_desc not in match_cache:
            vec = engine.embed_text(product_desc[:800])
            match_cache[product_desc] = resolve_tiers(vec)
        match = match_cache[product_desc]

        rec_item = match["tiers"].get(recommended_tier)
        if not rec_item:
            for fallback in ["best", "better", "good"]:
                if match["tiers"].get(fallback):
                    rec_item = match["tiers"][fallback]
                    break

        if not rec_item:
            line_rows.append(
                {
                    "month": month_label,
                    "location": location,
                    "product": product_desc,
                    "quantity": quantity,
                    "current_spend": current_spend,
                    "matched_category": None,
                    "recommended_item": None,
                    "recommended_spend": current_spend,
                    "savings": 0.0,
                    "matched": False,
                    "needs_manual_lookup": True,
                }
            )
            continue

        rec_uses_per_case = rec_item.get("uses_per_case")
        if comp_uses_per_case and rec_uses_per_case:
            total_uses = quantity * comp_uses_per_case
            recommended_cases = total_uses / rec_uses_per_case
            recommended_spend = recommended_cases * rec_item["price"]
        else:
            # No usage normalization data available: assume the same purchase
            # volume (case-for-case substitution). Directional, not exact.
            recommended_spend = quantity * rec_item["price"]

        line_rows.append(
            {
                "month": month_label,
                "location": location,
                "product": product_desc,
                "quantity": quantity,
                "current_spend": current_spend,
                "matched_category": match["matched_category"],
                "recommended_item": f"{rec_item['name']} ({rec_item['sku']})",
                "recommended_spend": recommended_spend,
                "savings": current_spend - recommended_spend,
                "matched": True,
                "needs_manual_lookup": False,
            }
        )

    detail_df = pd.DataFrame(line_rows)
    rows_unmatched = int((~detail_df["matched"]).sum())

    total_current = float(detail_df["current_spend"].sum())
    total_recommended = float(detail_df["recommended_spend"].sum())
    total_savings = total_current - total_recommended
    savings_pct = (total_savings / total_current * 100) if total_current else 0.0

    by_month = (
        detail_df.groupby("month")
        .agg(current_spend=("current_spend", "sum"), recommended_spend=("recommended_spend", "sum"))
        .reset_index()
    )
    by_month["savings"] = by_month["current_spend"] - by_month["recommended_spend"]
    by_month = by_month.sort_values("month")

    by_location = (
        detail_df.groupby("location")
        .agg(current_spend=("current_spend", "sum"), recommended_spend=("recommended_spend", "sum"))
        .reset_index()
    )
    by_location["savings"] = by_location["current_spend"] - by_location["recommended_spend"]
    by_location = by_location.sort_values("savings", ascending=False)

    matched_df = detail_df[detail_df["matched"]]
    by_category = (
        matched_df.groupby("matched_category")
        .agg(current_spend=("current_spend", "sum"), recommended_spend=("recommended_spend", "sum"))
        .reset_index()
    )
    by_category["savings"] = by_category["current_spend"] - by_category["recommended_spend"]
    by_category = by_category.sort_values("savings", ascending=False)

    return {
        "recommended_tier": recommended_tier,
        "total_current_spend": total_current,
        "total_recommended_spend": total_recommended,
        "total_savings": total_savings,
        "savings_pct": savings_pct,
        "rows_analyzed": len(detail_df),
        "rows_unmatched": rows_unmatched,
        "manual_lookup_note": (
            f"{rows_unmatched} record(s) had no catalog match. supplies.officebasics.com "
            "requires an account login and can't be searched automatically - check those "
            "products manually there until more catalog items are uploaded."
            if rows_unmatched
            else None
        ),
        "officebasics_portal_url": OFFICEBASICS_PORTAL_URL if rows_unmatched else None,
        "by_month": by_month.to_dict(orient="records"),
        "by_location": by_location.to_dict(orient="records"),
        "by_category": by_category.to_dict(orient="records"),
        "detail": detail_df.drop(columns=["matched"]).to_dict(orient="records"),
    }


def build_spend_analysis_workbook(analysis: Dict[str, Any]) -> str:
    money_fmt = MONEY_FMT  # local alias kept for readability in this function
    wb = openpyxl.Workbook()

    # --- Summary ---
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "Office Basics Spend Analysis"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = (
        f"Generated {datetime.now().strftime('%B %d, %Y')} \u2014 "
        f"recommended tier: {analysis['recommended_tier'].upper()}"
    )
    ws["A2"].font = NOTE_FONT

    summary_rows = [
        ("Total Current Spend", analysis["total_current_spend"], money_fmt),
        ("Total Recommended Spend", analysis["total_recommended_spend"], money_fmt),
        ("Total Savings", analysis["total_savings"], money_fmt),
        ("Savings %", analysis["savings_pct"] / 100, "0.0%"),
        ("Records Analyzed", analysis["rows_analyzed"], None),
        ("Records Unmatched", analysis["rows_unmatched"], None),
    ]
    for i, (label, value, fmt) in enumerate(summary_rows, start=4):
        ws.cell(row=i, column=1, value=label).font = Font(bold=True)
        c = ws.cell(row=i, column=2, value=value)
        if fmt:
            c.number_format = fmt
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 20

    # --- By Month ---
    ws2 = wb.create_sheet("By Month")
    headers = ["Month", "Current Spend", "Recommended Spend", "Savings"]
    for col, h in enumerate(headers, start=1):
        ws2.cell(row=1, column=col, value=h)
    style_header_row(ws2, 1, len(headers))
    for i, row in enumerate(analysis["by_month"], start=2):
        ws2.cell(row=i, column=1, value=row["month"])
        ws2.cell(row=i, column=2, value=row["current_spend"]).number_format = money_fmt
        ws2.cell(row=i, column=3, value=row["recommended_spend"]).number_format = money_fmt
        ws2.cell(row=i, column=4, value=row["savings"]).number_format = money_fmt
    for col, w in zip("ABCD", [14, 18, 20, 14]):
        ws2.column_dimensions[col].width = w
    if analysis["by_month"]:
        n = len(analysis["by_month"])
        chart = BarChart()
        chart.title = "Current vs. Recommended Spend by Month"
        chart.y_axis.title = "Spend ($)"
        data = Reference(ws2, min_col=2, max_col=3, min_row=1, max_row=n + 1)
        cats = Reference(ws2, min_col=1, min_row=2, max_row=n + 1)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        ws2.add_chart(chart, "F2")

    # --- By Location ---
    ws3 = wb.create_sheet("By Location")
    headers_loc = ["Location", "Current Spend", "Recommended Spend", "Savings"]
    for col, h in enumerate(headers_loc, start=1):
        ws3.cell(row=1, column=col, value=h)
    style_header_row(ws3, 1, len(headers_loc))
    for i, row in enumerate(analysis["by_location"], start=2):
        ws3.cell(row=i, column=1, value=row["location"])
        ws3.cell(row=i, column=2, value=row["current_spend"]).number_format = money_fmt
        ws3.cell(row=i, column=3, value=row["recommended_spend"]).number_format = money_fmt
        ws3.cell(row=i, column=4, value=row["savings"]).number_format = money_fmt
    for col, w in zip("ABCD", [22, 18, 20, 14]):
        ws3.column_dimensions[col].width = w
    if analysis["by_location"]:
        n = len(analysis["by_location"])
        chart2 = BarChart()
        chart2.title = "Savings Opportunity by Location"
        chart2.y_axis.title = "Spend ($)"
        data = Reference(ws3, min_col=2, max_col=3, min_row=1, max_row=n + 1)
        cats = Reference(ws3, min_col=1, min_row=2, max_row=n + 1)
        chart2.add_data(data, titles_from_data=True)
        chart2.set_categories(cats)
        ws3.add_chart(chart2, "F2")

    # --- By Category ---
    ws4 = wb.create_sheet("By Category")
    headers_cat = ["Category", "Current Spend", "Recommended Spend", "Savings"]
    for col, h in enumerate(headers_cat, start=1):
        ws4.cell(row=1, column=col, value=h)
    style_header_row(ws4, 1, len(headers_cat))
    for i, row in enumerate(analysis["by_category"], start=2):
        ws4.cell(row=i, column=1, value=row["matched_category"])
        ws4.cell(row=i, column=2, value=row["current_spend"]).number_format = money_fmt
        ws4.cell(row=i, column=3, value=row["recommended_spend"]).number_format = money_fmt
        ws4.cell(row=i, column=4, value=row["savings"]).number_format = money_fmt
    for col, w in zip("ABCD", [26, 18, 20, 14]):
        ws4.column_dimensions[col].width = w

    # --- Line Detail ---
    ws5 = wb.create_sheet("Line Detail")
    headers_d = [
        "Month", "Location", "Product", "Qty", "Current Spend",
        "Matched Category", "Recommended Item", "Recommended Spend", "Savings",
    ]
    for col, h in enumerate(headers_d, start=1):
        ws5.cell(row=1, column=col, value=h)
    style_header_row(ws5, 1, len(headers_d))
    for i, row in enumerate(analysis["detail"], start=2):
        ws5.cell(row=i, column=1, value=row["month"])
        ws5.cell(row=i, column=2, value=row["location"])
        ws5.cell(row=i, column=3, value=row["product"])
        ws5.cell(row=i, column=4, value=row["quantity"])
        ws5.cell(row=i, column=5, value=row["current_spend"]).number_format = money_fmt
        ws5.cell(row=i, column=6, value=row["matched_category"])
        ws5.cell(row=i, column=7, value=row["recommended_item"])
        ws5.cell(row=i, column=8, value=row["recommended_spend"]).number_format = money_fmt
        ws5.cell(row=i, column=9, value=row["savings"]).number_format = money_fmt
    for col_letter, w in zip("ABCDEFGHI", [12, 18, 30, 8, 14, 20, 30, 16, 12]):
        ws5.column_dimensions[col_letter].width = w

    filename = f"spend_analysis_{int(datetime.now().timestamp())}.xlsx"
    wb.save(os.path.join(REPORTS_DIR, filename))
    return filename


def build_executive_summary_docx(analysis: Dict[str, Any]) -> str:
    doc = Document()

    doc.add_heading("Office Basics Spend Analysis \u2014 Executive Summary", level=0)
    sub = doc.add_paragraph(f"Generated {datetime.now().strftime('%B %d, %Y')}")
    sub.runs[0].italic = True

    n_months = len(analysis["by_month"])
    n_locations = len(analysis["by_location"])
    doc.add_paragraph(
        f"This analysis reviewed {analysis['rows_analyzed']} purchase records across "
        f"{n_locations} location(s) and {n_months} month(s), totaling "
        f"${analysis['total_current_spend']:,.2f} in tracked spend. Based on "
        f"{analysis['recommended_tier'].upper()}-tier Office Basics equivalents, an estimated "
        f"${analysis['total_savings']:,.2f} ({analysis['savings_pct']:.1f}%) could have been saved "
        "over this period."
    )

    doc.add_heading("Key Figures", level=1)
    table = doc.add_table(rows=1, cols=2)
    table.style = "Light List Accent 1"
    hdr = table.rows[0].cells
    hdr[0].text, hdr[1].text = "Metric", "Value"
    for label, value in [
        ("Total Current Spend", f"${analysis['total_current_spend']:,.2f}"),
        ("Total Recommended Spend", f"${analysis['total_recommended_spend']:,.2f}"),
        ("Total Savings", f"${analysis['total_savings']:,.2f}"),
        ("Savings %", f"{analysis['savings_pct']:.1f}%"),
        ("Records Analyzed", str(analysis["rows_analyzed"])),
        ("Records Unmatched", str(analysis["rows_unmatched"])),
    ]:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value

    if analysis["by_category"]:
        doc.add_heading("Top Savings Opportunities by Category", level=1)
        top_categories = sorted(analysis["by_category"], key=lambda r: r["savings"], reverse=True)[:5]
        cat_table = doc.add_table(rows=1, cols=3)
        cat_table.style = "Light List Accent 1"
        hdr = cat_table.rows[0].cells
        hdr[0].text, hdr[1].text, hdr[2].text = "Category", "Current Spend", "Savings"
        for row in top_categories:
            cells = cat_table.add_row().cells
            cells[0].text = str(row["matched_category"])
            cells[1].text = f"${row['current_spend']:,.2f}"
            cells[2].text = f"${row['savings']:,.2f}"

    if analysis["by_location"]:
        doc.add_heading("Spend & Savings by Location", level=1)
        loc_table = doc.add_table(rows=1, cols=4)
        loc_table.style = "Light List Accent 1"
        hdr = loc_table.rows[0].cells
        hdr[0].text, hdr[1].text, hdr[2].text, hdr[3].text = (
            "Location", "Current Spend", "Recommended Spend", "Savings",
        )
        for row in analysis["by_location"]:
            cells = loc_table.add_row().cells
            cells[0].text = str(row["location"])
            cells[1].text = f"${row['current_spend']:,.2f}"
            cells[2].text = f"${row['recommended_spend']:,.2f}"
            cells[3].text = f"${row['savings']:,.2f}"

    doc.add_heading("Methodology & Assumptions", level=1)
    for bullet in [
        f"Recommended spend uses the {analysis['recommended_tier'].upper()} tier Office Basics "
        "equivalent for each matched product.",
        "Where per-case usage data (uses_per_case) is available for both the current product and "
        "its recommended replacement, spend is normalized to cost-per-use before comparison. "
        "Otherwise, a case-for-case substitution is assumed \u2014 treat those figures as directional.",
        f"{analysis['rows_unmatched']} of {analysis['rows_analyzed']} records could not be matched "
        "to a catalog category and are included in current spend totals with no assumed savings.",
        "Product matching uses a combination of previously confirmed recommendations and "
        "text/visual similarity against the loaded Office Basics catalog.",
    ]:
        doc.add_paragraph(bullet, style="List Bullet")

    if analysis.get("manual_lookup_note"):
        doc.add_paragraph(
            f"{analysis['manual_lookup_note']} Portal: {analysis['officebasics_portal_url']}",
            style="List Bullet",
        )

    doc.add_heading("Recommendation", level=1)
    doc.add_paragraph(
        "Prioritize rolling out Office Basics equivalents in the highest-savings categories and "
        "locations identified above, where the switching cost is lowest and the dollar impact is "
        "greatest."
    )

    filename = f"executive_summary_{int(datetime.now().timestamp())}.docx"
    doc.save(os.path.join(REPORTS_DIR, filename))
    return filename


@app.post("/analysis/spend/")
async def spend_analysis(
    file: UploadFile = File(...),
    recommended_tier: str = Form("best"),
):
    """
    Ingest historical spend data (date, location, product, quantity, and
    unit_price or total_price - optionally competitor_uses_per_case) and
    generate a usage analysis workbook plus an executive summary document
    showing how much switching to Office Basics equivalents would have saved.
    """
    if not engine.catalog:
        raise HTTPException(400, "No Office Basics catalog loaded. POST one to /catalog/upload/ first.")

    content = await file.read()
    ext = file.filename.lower().split(".")[-1]
    df = parse_spend_file(content, ext)
    analysis = analyze_spend(df, recommended_tier)

    xlsx_filename = build_spend_analysis_workbook(analysis)
    docx_filename = build_executive_summary_docx(analysis)

    return {
        "summary": {k: v for k, v in analysis.items() if k != "detail"},
        "workbook_url": f"/reports/{xlsx_filename}",
        "executive_summary_url": f"/reports/{docx_filename}",
    }
