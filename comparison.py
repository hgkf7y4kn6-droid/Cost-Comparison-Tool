"""The single-item cost-comparison Excel report."""
import os
from datetime import datetime
from typing import Any, Dict

import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from config import REPORTS_DIR
from excel_style import BORDER, MONEY_FMT, MONEY_FMT_PRECISE, NOTE_FONT, TITLE_FONT, style_header_row
from extraction import cost_per_use

def build_comparison_workbook(comparison: Dict[str, Any]) -> str:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Cost Comparison"

    ws["A1"] = "Office Basics Cost Comparison"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:F1")
    source_note = "Learned from a confirmed match" if comparison.get("source") == "learned" else "Algorithmic match"
    ws["A2"] = f"Generated {datetime.now().strftime('%B %d, %Y')} \u2014 {source_note}"
    ws["A2"].font = NOTE_FONT

    headers = ["Item", "SKU", "Description", "Price", "Cost / Use", "Savings vs Current"]
    header_row = 4
    for col, h in enumerate(headers, start=1):
        ws.cell(row=header_row, column=col, value=h)
    style_header_row(ws, header_row, len(headers))

    competitor_price = comparison.get("competitor_price")
    competitor_cpu = comparison.get("competitor_cost_per_use")

    rows = [
        {
            "label": f"CURRENT: {comparison['competitor_product']}",
            "sku": "\u2014",
            "description": "Currently purchased product (competitor)",
            "price": competitor_price,
            "cpu": competitor_cpu,
            "is_option": False,
        }
    ]
    for tier_name in ["good", "better", "best"]:
        item = comparison["tiers"].get(tier_name)
        if item:
            item_cpu = cost_per_use(item.get("price"), item.get("uses_per_case"))
            rows.append(
                {
                    "label": f"{tier_name.upper()}: {item['name']}",
                    "sku": item["sku"],
                    "description": item["description"],
                    "price": item["price"],
                    "cpu": item_cpu,
                    "is_option": True,
                }
            )

    baseline = competitor_cpu if competitor_cpu is not None else competitor_price

    r = header_row + 1
    for row in rows:
        ws.cell(row=r, column=1, value=row["label"]).border = BORDER
        ws.cell(row=r, column=2, value=row["sku"]).border = BORDER
        ws.cell(row=r, column=3, value=row["description"]).border = BORDER

        price_cell = ws.cell(row=r, column=4, value=row["price"])
        price_cell.number_format = MONEY_FMT
        price_cell.border = BORDER

        cpu_cell = ws.cell(row=r, column=5, value=row["cpu"])
        if row["cpu"] is not None:
            cpu_cell.number_format = MONEY_FMT_PRECISE
        else:
            cpu_cell.value = "n/a"
        cpu_cell.border = BORDER

        savings_cell = ws.cell(row=r, column=6)
        savings_cell.border = BORDER
        if row["is_option"] and baseline:
            compare_value = row["cpu"] if competitor_cpu is not None else row["price"]
            if compare_value is not None:
                savings = baseline - compare_value
                pct = (savings / baseline) * 100
                unit_label = "/use" if competitor_cpu is not None else ""
                savings_cell.value = f"${savings:,.4f}{unit_label} ({pct:.1f}%)"
                savings_cell.font = Font(color="1F7A1F" if savings > 0 else "B22222")
        r += 1

    for i, w in enumerate([34, 14, 45, 12, 12, 22], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    filename = f"cost_comparison_{int(datetime.now().timestamp())}.xlsx"
    wb.save(os.path.join(REPORTS_DIR, filename))
    return filename
